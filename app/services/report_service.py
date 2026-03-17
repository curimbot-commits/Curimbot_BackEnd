"""
Servicio para la generación de reportes en PDF y Excel.
app/services/report_service.py
"""

import io
import logging
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

logger = logging.getLogger(__name__)

class ReportService:
    @staticmethod
    def generate_pdf_report(user_name: str, stats_data: dict) -> bytes:
        """
        Genera un reporte PDF con las estadísticas semanales del usuario.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        
        # Estilos personalizados
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Title'],
            fontSize=18,
            textColor=colors.HexColor("#02ab74"),
            spaceAfter=20
        )
        
        elements = []
        
        # Título y Header
        elements.append(Paragraph(f"Resumen Semanal de Actividad", title_style))
        elements.append(Paragraph(f"Usuario: {user_name}", styles['Normal']))
        elements.append(Paragraph(f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Resumen General
        elements.append(Paragraph("Estadísticas Generales", styles['Heading2']))
        data = [
            ["Métrica", "Valor"],
            ["Total de Documentos", stats_data.get('totalDocuments', 0)],
            ["Documentos Hoy", stats_data.get('documentsToday', 0)],
            ["Documentos esta Semana", stats_data.get('documentsThisWeek', 0)],
            ["Tamaño Total (Bytes)", stats_data.get('totalSize', 0)],
            ["Promedio por Día", stats_data.get('averagePerDay', 0)],
        ]
        
        t = Table(data, colWidths=[200, 100])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#02ab74")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        elements.append(t)
        elements.append(Spacer(1, 20))
        
        # Desglose por tipo
        if stats_data.get('typeBreakdown'):
            elements.append(Paragraph("Desglose por Tipo de Archivo", styles['Heading2']))
            type_data = [["Tipo", "Cantidad", "Tamaño (Bytes)"]]
            for item in stats_data['typeBreakdown']:
                type_data.append([
                    item.get('file_type', 'unknown'),
                    item.get('count', 0),
                    item.get('size', 0)
                ])
            
            t_type = Table(type_data, colWidths=[100, 100, 150])
            t_type.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#7209b7")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey)
            ]))
            elements.append(t_type)
        
        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_excel_report(user_name: str, stats_data: dict) -> bytes:
        """
        Genera un reporte Excel con las estadísticas semanales del usuario.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumen Semanal"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="02ab74", end_color="02ab74", fill_type="solid")
        center_align = Alignment(horizontal="center")
        
        # Header
        ws['A1'] = "Resumen Semanal de Actividad"
        ws['A1'].font = Font(size=14, bold=True)
        ws['A2'] = f"Usuario: {user_name}"
        ws['A3'] = f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        
        # Tabla de Estadísticas Generales
        ws.append([]) # Fila vacía
        headers = ["Métrica", "Valor"]
        ws.append(headers)
        
        # Aplicar estilos a headers
        for cell in ws[ws.max_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            
        ws.append(["Total de Documentos", stats_data.get('totalDocuments', 0)])
        ws.append(["Documentos Hoy", stats_data.get('documentsToday', 0)])
        ws.append(["Documentos esta Semana", stats_data.get('documentsThisWeek', 0)])
        ws.append(["Tamaño Total (Bytes)", stats_data.get('totalSize', 0)])
        ws.append(["Promedio por Día", stats_data.get('averagePerDay', 0)])
        
        # Desglose por tipo
        if stats_data.get('typeBreakdown'):
            ws.append([])
            ws.append(["Desglose por Tipo de Archivo"])
            ws[ws.max_row][0].font = Font(bold=True)
            
            type_headers = ["Tipo", "Cantidad", "Tamaño (Bytes)"]
            ws.append(type_headers)
            
            # Estilos para type_headers
            type_fill = PatternFill(start_color="7209b7", end_color="7209b7", fill_type="solid")
            for cell in ws[ws.max_row]:
                cell.font = header_font
                cell.fill = type_fill
                cell.alignment = center_align
                
            for item in stats_data['typeBreakdown']:
                ws.append([
                    item.get('file_type', 'unknown'),
                    item.get('count', 0),
                    item.get('size', 0)
                ])
        
        # Ajustar ancho de columnas
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes
